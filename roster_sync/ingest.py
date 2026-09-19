"""Reading an agency roster workbook into normalized rows.

Agency spreadsheets are written for humans, not for parsers. They carry a
title row above the headers, blank spacer rows, trailing notes below the
data, and columns whose headers differ week to week. This module locates the
header row rather than assuming row 1, maps columns through header aliases
rather than position, and reports what it could not understand instead of
failing on the first bad cell. A header it cannot read is the exception: that
raises, because a file with no name or contact column rejects every row and
still counts as a processed period.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .models import RosterRow
from .normalize import (
    PLACEHOLDER_TOKENS,
    normalize_email,
    normalize_name,
    normalize_phone,
    normalize_role,
)

# Canonical field -> header spellings seen in the wild. Extending this is a
# one-line change here, or pass header_aliases to read_roster; nothing in the
# parsing logic knows about column order.
DEFAULT_HEADER_ALIASES: dict[str, list[str]] = {
    "first_name": ["first name", "first", "firstname", "given name", "fname"],
    "last_name": ["last name", "last", "lastname", "surname", "lname"],
    "phone": ["phone", "phone number", "mobile", "cell", "cell phone", "contact number"],
    "email": ["email", "email address", "e-mail", "mail"],
    "role": ["role", "position", "job title", "title", "assignment", "job"],
}

REQUIRED_FIELDS = ("first_name", "last_name")


@dataclass
class IngestReport:
    """What the parser found, for logging and for the exception report."""

    path: str
    sheet: str
    header_row: int
    column_map: dict[str, int]
    rows_read: int
    rows_blank: int
    # (source_row, text as written) for each role cell the role map could not
    # read. Reported for this run only; the text is never stored.
    unmapped_roles: list[tuple[int, str]]


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip().lower()


def find_header_row(
    grid: list[tuple],
    aliases: dict[str, list[str]],
    scan_depth: int = 15,
) -> tuple[int, dict[str, int]]:
    """Locate the header row and map canonical fields to column indices.

    Scores each of the first scan_depth rows by how many recognized headers
    it contains, and takes the best. A title row scores zero; the real header
    row scores highest.
    """
    lookup = {s: f for f, names in aliases.items() for s in names}
    best_row, best_map, best_score = -1, {}, 0

    for index, row in enumerate(grid[:scan_depth]):
        column_map: dict[str, int] = {}
        for column, value in enumerate(row):
            field = lookup.get(_cell_text(value))
            if field and field not in column_map:
                column_map[field] = column
        if len(column_map) > best_score:
            best_row, best_map, best_score = index, column_map, len(column_map)

    if best_score == 0:
        raise ValueError(
            "no recognizable header row in the first "
            f"{scan_depth} rows; check the file or extend the header aliases"
        )
    return best_row, best_map


def read_roster(
    path: str | Path,
    role_map: dict[str, str],
    header_aliases: dict[str, list[str]] | None = None,
    sheet: str | None = None,
) -> tuple[list[RosterRow], IngestReport]:
    """Parse a roster workbook into normalized rows plus an ingest report."""
    aliases = header_aliases or DEFAULT_HEADER_ALIASES
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]

    # Everything needed is read before anything can raise, so a rejected file
    # does not leak the read-only handle.
    grid = list(worksheet.iter_rows(values_only=True))
    sheet_title = worksheet.title
    workbook.close()

    header_row, column_map = find_header_row(grid, aliases)

    # Fail closed: without these columns no row is usable, the period is still
    # recorded, and the second such week ages every active worker into a leaver.
    unmapped = [f for f in REQUIRED_FIELDS if f not in column_map]
    if "phone" not in column_map and "email" not in column_map:
        unmapped.append("phone or email")
    if unmapped:
        found = [str(cell).strip() for cell in grid[header_row] if _cell_text(cell)]
        raise ValueError(
            f"header row {header_row + 1} has no column for {', '.join(unmapped)}; "
            f"headers found: {found}; extend the header aliases or pass header_aliases"
        )

    rows: list[RosterRow] = []
    unmapped_roles: list[tuple[int, str]] = []
    blank = 0

    for offset, raw_row in enumerate(grid[header_row + 1:], start=header_row + 2):
        if all(cell is None or str(cell).strip() == "" for cell in raw_row):
            blank += 1
            continue

        def value(field: str) -> object:
            column = column_map.get(field)
            if column is None or column >= len(raw_row):
                return None
            return raw_row[column]

        role_cell = value("role")
        role = normalize_role(role_cell, role_map)
        # A placeholder counts as an empty cell, as it does in normalize_role:
        # it names no role, and no alias in the role map would fix it.
        role_unmapped = role is None and _cell_text(role_cell) not in PLACEHOLDER_TOKENS
        if role_unmapped:
            unmapped_roles.append((offset, str(role_cell).strip()))

        rows.append(
            RosterRow(
                source_row=offset,
                name=normalize_name(value("first_name"), value("last_name")),
                phone=normalize_phone(value("phone")),
                email=normalize_email(value("email")),
                role=role,
                role_unmapped=role_unmapped,
            )
        )

    report = IngestReport(
        path=str(path),
        sheet=sheet_title,
        header_row=header_row + 1,
        column_map=column_map,
        rows_read=len(rows),
        rows_blank=blank,
        unmapped_roles=unmapped_roles,
    )
    return rows, report
